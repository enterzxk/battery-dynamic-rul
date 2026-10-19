import pathlib,json,hashlib,zipfile,io,time,signal,fcntl,math,re,statistics,traceback
import openpyxl
R=pathlib.Path('/home/yndx/zxk/rul');D=pathlib.Path(__file__).resolve().parent
sha=lambda b:hashlib.sha256(b).hexdigest()
P=json.loads((D/'protocol.json').read_text());G=json.loads((D/'ROOT_READ_GO.json').read_text())
assert G['protocol_sha256']==sha((D/'protocol.json').read_bytes())
assert G['runner_sha256']==sha(pathlib.Path(__file__).read_bytes())
lock=open(R/'automation/state/online_survival_cpu.lock','a');fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
C=json.loads((R/'automation/state/coordinator.json').read_text());assert C['owner']==G['owner'] and C['new_fits_allowed'] is False
with open(D/'claim.json','x') as f:json.dump({'status':'CONSUMED_BEFORE_MEMBER_READ','go_sha256':sha((D/'ROOT_READ_GO.json').read_bytes())},f)
start=time.monotonic();signal.signal(signal.SIGALRM,lambda *a:(_ for _ in ()).throw(TimeoutError('60s budget')));signal.alarm(P['max_seconds'])
out={'new_fits':0,'other_cell_reads':0,'other_date_payload_reads':0,'model_inference':0,'member_payload_reads':0}
try:
 b=(R/P['archive']).read_bytes();assert sha(b)==P['archive_sha256']
 with zipfile.ZipFile(io.BytesIO(b)) as z:
  i=z.getinfo(P['member']);assert i.file_size==P['member_bytes'] and format(i.CRC,'08x')==P['crc32']
  out['member_payload_reads']=1; raw=z.read(i)
 with zipfile.ZipFile(io.BytesIO(raw)) as z:
  assert len(z.infolist())<=P['max_xlsx_entries'] and sum(i.file_size for i in z.infolist())<=P['max_xlsx_uncompressed_bytes']
 out['member_sha256']=sha(raw)
 wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
 assert len(wb.worksheets)<=P['max_sheets'];sheets=[];matches=[]
 norm=lambda x:re.sub('[^a-z0-9]','',str(x).lower())
 required={'cycleindex','testtimes','currenta','voltagev','chargecapacityah','dischargecapacityah'}
 for ws in wb.worksheets:
  assert ws.max_column<=P['max_columns'] and ws.max_row<=P['max_rows_per_sheet']
  rows=list(ws.iter_rows(values_only=True));head=None
  for j,row in enumerate(rows[:10]):
   if required.issubset(set(map(norm,row))):
    assert head is None;head=j
  sheets.append({'name':ws.title,'rows':len(rows),'columns':ws.max_column,'first_rows':[[str(v) if v is not None else None for v in row] for row in rows[:3]],'header_row_1based':None if head is None else head+1})
  if head is not None:matches.append((ws.title,rows,head))
 wb.close();out['sheets']=sheets;assert len(matches)==1,'missing/ambiguous data schema'
 title,rows,h=matches[0];ix={norm(v):j for j,v in enumerate(rows[h])};cycles={};lastt=None
 for row in rows[h+1:]:
  if all(v is None for v in row):continue
  vals={k:float(row[ix[k]]) for k in required};assert all(math.isfinite(v) for v in vals.values());assert vals['cycleindex'].is_integer()
  t=vals['testtimes'];assert lastt is None or t>=lastt,'test time decreases';lastt=t
  c=int(vals['cycleindex']);cycles.setdefault(c,[]).append(vals)
 summary=[]
 for c,rr in cycles.items():
  neg=[(j,x) for j,x in enumerate(rr) if x['currenta']<-.001];pos=[(j,x) for j,x in enumerate(rr) if x['currenta']>.001]
  item={'source_cycle_index':c,'rows':len(rr),'negative_rows':len(neg),'positive_rows':len(pos)}
  if neg:
   first=neg[0][0];prev=[x for j,x in pos if j<first];last=neg[-1][1]
   integral=sum(max(0,-(a['currenta']+b['currenta'])/2)*(b['testtimes']-a['testtimes'])/3600 for a,b in zip(rr,rr[1:]) if a['currenta']<-.001 and b['currenta']<-.001)
   item.update(first_negative_time_s=neg[0][1]['testtimes'],last_negative_voltage_V=last['voltagev'],discharge_counter_Ah=last['dischargecapacityah'],negative_segment_integral_Ah=integral,negative_current_median_A=statistics.median(x['currenta'] for j,x in neg),prior_positive_sample=None if not prev else {k:prev[-1][k] for k in ['testtimes','currenta','voltagev','chargecapacityah']},counter_minus_integral_Ah=last['dischargecapacityah']-integral)
  summary.append(item)
 out.update(status='READ_COMPLETE_DESCRIPTIVE_ONLY',data_sheet=title,source_cycle_count=len(summary),cycles=summary,raw_data_publication=False,full_capacity_qualified=False,EOL_created=False,limitations='Recorded boundary samples are not exact termination proof. First file cannot establish lifetime, cycle continuity across dates, identity or generalization.')
except Exception as e:out.update(status='FAILED_STOP_NO_RETRY',error=type(e).__name__+': '+str(e))
finally:
 signal.alarm(0);out['elapsed_seconds']=time.monotonic()-start
 with open(D/'summary.json','x') as f:json.dump(out,f,indent=2,default=str)
 print(json.dumps(out,default=str))
