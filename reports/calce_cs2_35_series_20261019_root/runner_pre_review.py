import pathlib,json,hashlib,zipfile,io,time,signal,fcntl,math,re,statistics,csv,datetime
import openpyxl
R=pathlib.Path('/home/yndx/zxk/rul');D=pathlib.Path(__file__).resolve().parent;sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
P=json.loads((D/'protocol.json').read_text());G=json.loads((D/'ROOT_READ_GO.json').read_text());assert G['protocol_sha256']==sha(D/'protocol.json') and G['runner_sha256']==sha(pathlib.Path(__file__))
lock=open(R/'automation/state/online_survival_cpu.lock','a');fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB);C=json.loads((R/'automation/state/coordinator.json').read_text());assert C['owner']==G['owner'] and C['new_fits_allowed'] is False
with open(D/'claim.json','x') as f:json.dump({'consumed':True,'go_sha256':sha(D/'ROOT_READ_GO.json')},f)
start=time.monotonic();signal.signal(signal.SIGALRM,lambda *a:(_ for _ in ()).throw(TimeoutError('300s')));signal.alarm(P['max_seconds']);O={'status':'RUNNING','member_reads':0,'member_bytes':0,'data_rows':0,'files':[],'new_fits':0,'model_inference':0,'other_cell_reads':0,'firstfile_payload_reads':0,'EOL_created':False};allcycles=[]
try:
 F=R/'reports/calce_cs2_35_firstfile_20261019_root/summary.json';assert sha(F)==P['firstfile_saved_summary_sha256'];first=json.loads(F.read_text());O['previous_firstfile_summary_reused']=True
 assert len(P['members'])==24 and sha(R/P['archive'])==P['archive_sha256']
 norm=lambda x:re.sub('[^a-z0-9]','',str(x).lower());required={'cycleindex','testtimes','currenta','voltagev','chargecapacityah','dischargecapacityah','datetime'}
 with zipfile.ZipFile(R/P['archive']) as z:
  for entry in P['members']:
   assert entry['name']!='CS2_35/CS2_35_8_17_10.xlsx'
   i=z.getinfo(entry['name']);assert i.file_size==entry['bytes']<=P['max_member_bytes'] and format(i.CRC,'08x')==entry['crc32'];assert O['member_bytes']+i.file_size<=P['max_total_member_bytes']
   O['member_reads']+=1;raw=z.read(i);O['member_bytes']+=len(raw)
   with zipfile.ZipFile(io.BytesIO(raw)) as inner:assert len(inner.infolist())<=P['max_xlsx_entries'] and sum(j.file_size for j in inner.infolist())<=P['max_xlsx_internal_bytes']
   wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False);assert len(wb.worksheets)<=P['max_sheets'];matches=[]
   for ws in wb.worksheets:
    assert ws.max_column<=P['max_columns'] and ws.max_row<=P['max_rows_per_sheet'];rows=list(ws.iter_rows(values_only=True))
    for j,row in enumerate(rows[:10]):
     if required.issubset(set(map(norm,row))):matches.append((ws.title,rows,j))
   wb.close();assert len(matches)==1,'schema ambiguous/missing';title,rows,h=matches[0];ix={norm(v):j for j,v in enumerate(rows[h])};groups=[];lastt=None;firsttime=None;lastdate=None;firstdate=None;file_rows=0
   for row in rows[h+1:]:
    if all(v is None for v in row):continue
    v={k:float(row[ix[k]]) for k in required-{'datetime'}};assert all(math.isfinite(x) for x in v.values());assert v['cycleindex'].is_integer();t=v['testtimes'];assert lastt is None or t>=lastt,'within-file test time reset';lastt=t
    dt=row[ix['datetime']];assert isinstance(dt,datetime.datetime),'Date_Time not parsed datetime';v['actual_datetime']=dt.isoformat();lastdate=dt.isoformat()
    if firsttime is None:firsttime=t;firstdate=lastdate
    c=int(v['cycleindex'])
    if not groups or groups[-1][0]!=c:groups.append((c,[]))
    groups[-1][1].append(v);file_rows+=1;O['data_rows']+=1;assert O['data_rows']<=P['max_total_rows']
   for occ,(c,rr) in enumerate(groups):
    neg=[(j,v) for j,v in enumerate(rr) if v['currenta']<-.001];pos=[(j,v) for j,v in enumerate(rr) if v['currenta']>.001];a={'file':entry['name'],'occurrence':occ,'source_cycle':c,'rows':len(rr),'start_datetime':rr[0]['actual_datetime'],'end_datetime':rr[-1]['actual_datetime'],'negative_rows':len(neg),'positive_rows':len(pos)}
    if neg:
     prev=[v for j,v in pos if j<neg[0][0]];last=neg[-1][1];integral=sum(-(u['currenta']+v['currenta'])/2*(v['testtimes']-u['testtimes'])/3600 for u,v in zip(rr,rr[1:]) if u['currenta']<-.001 and v['currenta']<-.001)
     a.update(discharge_counter_Ah=last['dischargecapacityah'],partial_integral_Ah=integral,end_discharge_V=last['voltagev'],end_discharge_A=last['currenta'],prior_charge_V=None if not prev else prev[-1]['voltagev'],prior_charge_A=None if not prev else prev[-1]['currenta'],counter_minus_partial_integral_Ah=last['dischargecapacityah']-integral)
    allcycles.append(a)
   assert groups,'empty data';O['files'].append({'member':entry['name'],'sha256':hashlib.sha256(raw).hexdigest(),'bytes':len(raw),'sheet':title,'rows':file_rows,'occurrences':len(groups),'first_source_cycle':groups[0][0],'last_source_cycle':groups[-1][0],'first_test_time_s':firsttime,'last_test_time_s':lastt,'first_actual_datetime':firstdate,'last_actual_datetime':lastdate})
 O['boundaries']=[{'previous':a['member'],'next':b['member'],'cycle_change':b['first_source_cycle']-a['last_source_cycle'],'test_time_change_s':b['first_test_time_s']-a['last_test_time_s'],'wall_time_gap_s':(datetime.datetime.fromisoformat(b['first_actual_datetime'])-datetime.datetime.fromisoformat(a['last_actual_datetime'])).total_seconds()} for a,b in zip(O['files'],O['files'][1:])]
 vals=[x['discharge_counter_Ah'] for x in allcycles if 'discharge_counter_Ah' in x];O.update(status='COMPLETE_DESCRIPTIVE_SERIES_NOT_EOL',occurrences=len(allcycles),discharging_occurrences=len(vals),capacity_counter_min_Ah=min(vals),capacity_counter_max_Ah=max(vals),source_cycle_min=min(x['source_cycle'] for x in allcycles),source_cycle_max=max(x['source_cycle'] for x in allcycles),missing_prior_charge=sum(x['negative_rows']>0 and x.get('prior_charge_V') is None for x in allcycles),negative_boundary_walltime_gaps=sum(x['wall_time_gap_s']<0 for x in O['boundaries']),limitations='No cross-file merge, native cycle continuity/complete capacity/EOL qualification or independent identity claim. All sampled occurrences retained; canonical table server only.')
except Exception as e:O.update(status='FAILED_STOP_NO_RETRY',error=type(e).__name__+': '+str(e))
finally:
 signal.alarm(0);O['elapsed_seconds']=time.monotonic()-start
 keys=sorted(set().union(*(x.keys() for x in allcycles))) if allcycles else ['file']
 with open(D/'canonical_occurrences.csv','x',newline='') as f:
  w=csv.DictWriter(f,fieldnames=keys);w.writeheader();w.writerows(allcycles)
 O['canonical_sha256']=sha(D/'canonical_occurrences.csv');O['saved_occurrences']=len(allcycles)
 with open(D/'summary.json','x') as f:json.dump(O,f,indent=2)
 print(json.dumps({k:v for k,v in O.items() if k not in ['files','boundaries']}))
